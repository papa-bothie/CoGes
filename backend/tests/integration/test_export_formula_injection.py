"""Aucune formule dans un export, et l'aller-retour export → import reste fidèle (ADR 0045).

Un tiers (client, fournisseur, vendeur) peut saisir `=HYPERLINK(…)` comme nom : ce texte
ne doit jamais devenir une formule active chez le gérant qui ouvre l'export.
"""

from __future__ import annotations

import csv
import io
import zipfile
from typing import Any
from uuid import uuid4

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.core import queue
from app.workers.tasks import exports as export_tasks
from tests.conftest import (
    _create_demo_company,
    auth_headers,
    login_token,
    write_headers,
)
from tests.integration.test_exports import _request


@pytest.fixture
def enqueued(monkeypatch: pytest.MonkeyPatch) -> list[tuple[str, dict[str, Any]]]:
    """L'enfilement est intercepté ; la tâche est exécutée en direct sur la base de test."""
    calls: list[tuple[str, dict[str, Any]]] = []

    def record(task_name: str, *, countdown: int = 0, **kwargs: Any) -> bool:
        calls.append((task_name, kwargs))
        return True

    monkeypatch.setattr(queue, "enqueue", record)
    return calls


def _names(marker: str) -> list[str]:
    return [
        '=HYPERLINK("http://evil.example","' + marker + '")',
        f"+221 {marker}",
        f"-2+3 {marker}",
        f"@SUM(A1) {marker}",
    ]


def _visible_rows(content: bytes, fmt: str) -> list[list[Any]]:
    """Ce que le tableur affiche, après avoir vérifié qu'il n'y a aucune formule."""
    if fmt == "csv":
        return list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert b"<f>" not in zipfile.ZipFile(io.BytesIO(content)).read("xl/worksheets/sheet1.xml")
    sheet = load_workbook(io.BytesIO(content)).worksheets[0]
    rows = list(sheet.iter_rows())
    assert all(cell.data_type != "f" for row in rows for cell in row)
    return [[cell.value for cell in row] for row in rows]


def _expected(name: str, fmt: str) -> str:
    """En CSV, l'apostrophe est le seul moyen de neutraliser ; en .xlsx, rien n'est ajouté."""
    return "'" + name if fmt == "csv" else name


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_customer_names_never_become_formulas(
    api_client: TestClient, alpha_token: str, fmt: str
) -> None:
    names = _names(uuid4().hex[:8])
    for name in names:
        created = api_client.post(
            "/api/v1/customers", headers=auth_headers(alpha_token), json={"name": name}
        )
        assert created.status_code == 201, created.text

    exported = api_client.get(
        "/api/v1/exports/referentials/customers",
        params={"format": fmt},
        headers=auth_headers(alpha_token),
    )
    assert exported.status_code == 200, exported.text

    shown = {row[0] for row in _visible_rows(exported.content, fmt)}
    assert {_expected(name, fmt) for name in names} <= shown


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_product_names_never_become_formulas(
    api_client: TestClient, alpha_token: str, fmt: str
) -> None:
    marker = uuid4().hex[:8]
    names = _names(marker)
    for index, name in enumerate(names):
        created = api_client.post(
            "/api/v1/products",
            headers=auth_headers(alpha_token),
            json={"sku": f"FX-{marker}-{index}", "name": name, "unit_price": "1500"},
        )
        assert created.status_code == 201, created.text

    exported = api_client.get(
        "/api/v1/exports/referentials/products",
        params={"format": fmt},
        headers=auth_headers(alpha_token),
    )
    assert exported.status_code == 200, exported.text

    rows = _visible_rows(exported.content, fmt)
    header = rows[0]
    ours = {row[header.index("nom")]: row for row in rows[1:] if str(row[0]).startswith("FX-")}
    assert {_expected(name, fmt) for name in names} <= set(ours)
    if fmt == "xlsx":
        # Le prix reste un nombre dans le classeur.
        prices = [row[header.index("prix_vente")] for row in ours.values()]
        assert all(isinstance(price, int | float) for price in prices)


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_sales_export_never_holds_a_formula(
    api_client: TestClient,
    alpha_token: str,
    workers_bound_to_test_database: None,
    enqueued: list[tuple[str, dict[str, Any]]],
    fmt: str,
) -> None:
    marker = uuid4().hex[:8]
    product = api_client.post(
        "/api/v1/products",
        headers=auth_headers(alpha_token),
        json={"sku": f"FXV-{marker}", "name": f"-Article {marker}", "unit_price": "1500"},
    ).json()["id"]
    stocked = api_client.post(
        "/api/v1/inventory/movements",
        headers=auth_headers(alpha_token),
        json={"product_id": product, "type": "initial", "qty_delta": "10", "unit_cost": "1000"},
    )
    assert stocked.status_code == 201, stocked.text

    names = _names(marker)
    for name in names:
        customer = api_client.post(
            "/api/v1/customers", headers=auth_headers(alpha_token), json={"name": name}
        ).json()["id"]
        sale = api_client.post(
            "/api/v1/sales", headers=auth_headers(alpha_token), json={"customer_id": customer}
        ).json()["id"]
        line = api_client.post(
            f"/api/v1/sales/{sale}/items",
            headers=auth_headers(alpha_token),
            json={"product_id": product, "qty": "1"},
        )
        assert line.status_code == 201, line.text
        confirmed = api_client.post(
            f"/api/v1/sales/{sale}/confirm",
            headers=auth_headers(alpha_token),
            json={"paid_amount": "1500.00", "method": "cash"},
        )
        assert confirmed.status_code == 200, confirmed.text

    requested = _request(api_client, alpha_token, format=fmt)
    assert requested.status_code == 202, requested.text
    [kwargs] = [
        call_kwargs
        for task_name, call_kwargs in enqueued
        if task_name == "exports.build_sales" and call_kwargs["export_id"] == requested.json()["id"]
    ]
    assert export_tasks.build_sales(**kwargs) == "ready"
    downloaded = api_client.get(
        f"/api/v1/exports/{requested.json()['id']}/file", headers=auth_headers(alpha_token)
    )
    assert downloaded.status_code == 200

    rows = _visible_rows(downloaded.content, fmt)
    header = rows[0]
    shown = {row[header.index("client")] for row in rows[1:]}
    assert {_expected(name, fmt) for name in names} <= shown


@pytest.mark.parametrize("fmt", ["csv", "xlsx"])
def test_export_then_import_gives_back_the_same_customers(api_client: TestClient, fmt: str) -> None:
    """Deux entreprises neuves : l'export de l'une, réimporté chez l'autre, se réexporte
    à l'identique — apostrophes comprises."""
    marker = uuid4().hex[:8]
    source = _create_demo_company(f"fx-source-{fmt}-{marker}")
    target = _create_demo_company(f"fx-cible-{fmt}-{marker}")
    source_token = login_token(api_client, source)
    target_token = login_token(api_client, target)

    for index, name in enumerate(_names(marker)):
        created = api_client.post(
            "/api/v1/customers",
            headers=auth_headers(source_token),
            json={"name": name, "code": f"FX-{index}", "phone": f"+221 77 000 00 0{index}"},
        )
        assert created.status_code == 201, created.text

    def export(token: str) -> bytes:
        response = api_client.get(
            "/api/v1/exports/referentials/customers",
            params={"format": fmt},
            headers=auth_headers(token),
        )
        assert response.status_code == 200, response.text
        return response.content

    exported = export(source_token)
    imported = api_client.post(
        "/api/v1/imports/customers",
        headers=write_headers(target_token),
        files={"file": (f"clients.{fmt}", exported, "application/octet-stream")},
    )
    assert imported.status_code == 201, imported.text

    stored = api_client.get("/api/v1/customers?page_size=50", headers=auth_headers(target_token))
    assert sorted(item["name"] for item in stored.json()["items"]) == sorted(_names(marker))
    assert all(item["phone"].startswith("+221") for item in stored.json()["items"])
    assert _visible_rows(export(target_token), fmt) == _visible_rows(exported, fmt)

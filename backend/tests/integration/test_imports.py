"""Import des référentiels (phase 10, ADR 0041) — contre la vraie base, en `coges_app`.

Ce que le module promet, et que ces tests vérifient **en base** :

* un fichier fautif n'écrit **rien** — ni produit, ni prix, ni mouvement, ni audit,
  ni événement — et le rapport nomme chaque ligne et chaque colonne fautives ;
* la simulation (`dry_run`) rend le même rapport que l'import réel, sans écrire ;
* le stock initial passe par un mouvement `initial` et fixe le CMP ;
* un rejeu avec la même clé ne crée rien de plus.
"""

from __future__ import annotations

import io
from typing import Any
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import sessionmaker

from tests.conftest import DemoCompany, auth_headers, read_scalar_in_tenant, write_headers

PRODUCT_HEADER = (
    "sku;nom;prix_vente;categorie;unite;taux_taxe;prix_achat;stock_initial;cout_unitaire"
)


def _upload(
    client: TestClient,
    token: str,
    kind: str,
    content: bytes,
    *,
    dry_run: bool = False,
    key: str | None = None,
    filename: str = "import.csv",
) -> Any:
    return client.post(
        f"/api/v1/imports/{kind}",
        params={"dry_run": "true"} if dry_run else None,
        headers=write_headers(token, key),
        files={"file": (filename, content, "application/octet-stream")},
    )


def _count(
    session_factory: sessionmaker[Any], company: DemoCompany, sql: str, **params: Any
) -> int:
    return int(read_scalar_in_tenant(session_factory, company.company_id, sql, params))


def _products_csv(prefix: str, count: int, *, broken_row: int | None = None) -> bytes:
    lines = [PRODUCT_HEADER]
    for index in range(1, count + 1):
        price = "abc" if index == broken_row else f"{1000 + index} "
        lines.append(f"{prefix}-{index:02d};Produit {index};{price};;pièce;18 %;600;10;600,50")
    return "\n".join(lines).encode("utf-8")


def _state(session_factory: sessionmaker[Any], company: DemoCompany, prefix: str) -> dict[str, int]:
    like = f"{prefix}-%"
    return {
        "products": _count(
            session_factory, company, "SELECT count(*) FROM products WHERE sku LIKE :p", p=like
        ),
        "prices": _count(
            session_factory,
            company,
            "SELECT count(*) FROM product_prices pp JOIN products p ON p.id = pp.product_id "
            "WHERE p.sku LIKE :p",
            p=like,
        ),
        "movements": _count(
            session_factory,
            company,
            "SELECT count(*) FROM inventory_movements m JOIN products p ON p.id = m.product_id "
            "WHERE p.sku LIKE :p",
            p=like,
        ),
        "events": _count(session_factory, company, "SELECT count(*) FROM business_events"),
        "audit": _count(session_factory, company, "SELECT count(*) FROM audit_logs"),
    }


# ------------------------------------------------------------------ tout ou rien


def test_file_with_an_error_writes_nothing(
    api_client: TestClient,
    alpha_token: str,
    company_alpha: DemoCompany,
    session_factory: sessionmaker[Any],
) -> None:
    prefix = f"ERR{uuid4().hex[:6].upper()}"
    before = _state(session_factory, company_alpha, prefix)

    response = _upload(api_client, alpha_token, "products", _products_csv(prefix, 20, broken_row=7))

    assert response.status_code == 422, response.text
    error = response.json()["error"]
    assert error["code"] == "import_invalid"
    report = error["details"]
    assert report["rows"] == 20
    assert report["created"] == 0
    assert [(item["row"], item["column"]) for item in report["errors"]] == [(8, "Prix de vente")]
    # Rien n'a été écrit : ni référentiel, ni stock, ni audit, ni événement.
    assert _state(session_factory, company_alpha, prefix) == before


def test_dry_run_reports_without_writing(
    api_client: TestClient,
    alpha_token: str,
    company_alpha: DemoCompany,
    session_factory: sessionmaker[Any],
) -> None:
    prefix = f"DRY{uuid4().hex[:6].upper()}"
    before = _state(session_factory, company_alpha, prefix)

    clean = _upload(api_client, alpha_token, "products", _products_csv(prefix, 5), dry_run=True)
    broken = _upload(
        api_client, alpha_token, "products", _products_csv(prefix, 5, broken_row=2), dry_run=True
    )

    assert clean.status_code == 200, clean.text
    assert clean.json() == {
        "kind": "products",
        "dry_run": True,
        "rows": 5,
        "created": 5,
        "errors": [],
    }
    assert broken.status_code == 200
    assert [item["row"] for item in broken.json()["errors"]] == [3]
    assert _state(session_factory, company_alpha, prefix) == before


def test_twenty_products_with_one_error_then_corrected(
    api_client: TestClient,
    alpha_token: str,
    company_alpha: DemoCompany,
    session_factory: sessionmaker[Any],
) -> None:
    """La démonstration de la phase 10 : erreur volontaire, correction, import."""
    prefix = f"DEMO{uuid4().hex[:6].upper()}"
    refused = _upload(api_client, alpha_token, "products", _products_csv(prefix, 20, broken_row=13))
    assert refused.status_code == 422

    imported = _upload(api_client, alpha_token, "products", _products_csv(prefix, 20))

    assert imported.status_code == 201, imported.text
    assert imported.json()["created"] == 20
    state = _state(session_factory, company_alpha, prefix)
    assert state["products"] == 20
    assert state["prices"] == 20
    assert state["movements"] == 20
    # Le stock initial est un mouvement `initial`, et le CMP part du coût saisi.
    initial = _count(
        session_factory,
        company_alpha,
        "SELECT count(*) FROM inventory_movements m JOIN products p ON p.id = m.product_id "
        "WHERE p.sku LIKE :p AND m.type = 'initial' AND m.qty_delta = 10 AND m.unit_cost = 600.5",
        p=f"{prefix}-%",
    )
    assert initial == 20
    stock = read_scalar_in_tenant(
        session_factory,
        company_alpha.company_id,
        "SELECT s.avg_cost FROM product_stock s JOIN products p ON p.id = s.product_id "
        "WHERE p.sku = :sku",
        {"sku": f"{prefix}-01"},
    )
    assert str(stock) == "600.5000"
    # « 18 % » est devenu 0.18, « pièce » l'unité `piece`.
    listed = api_client.get(
        "/api/v1/products", params={"search": f"{prefix}-01"}, headers=auth_headers(alpha_token)
    ).json()["items"]
    product = next(item for item in listed if item["sku"] == f"{prefix}-01")
    assert product["default_tax_rate"] == "0.1800"
    assert product["unit"] == "piece"


# ------------------------------------------------------------------ doublons


def test_duplicate_sku_in_file_and_in_database(
    api_client: TestClient,
    alpha_token: str,
    company_alpha: DemoCompany,
    session_factory: sessionmaker[Any],
) -> None:
    prefix = f"DUP{uuid4().hex[:6].upper()}"
    assert _upload(api_client, alpha_token, "products", _products_csv(prefix, 2)).status_code == 201

    content = "\n".join(
        [
            "sku;nom;prix_vente",
            f"{prefix}-01;Déjà en base;1500",
            f"{prefix}-NEW;Nouveau;1500",
            f"{prefix}-NEW;Doublon du fichier;1500",
        ]
    ).encode()
    response = _upload(api_client, alpha_token, "products", content, dry_run=True)

    errors = {(item["row"], item["code"]) for item in response.json()["errors"]}
    assert errors == {(2, "product_sku_taken"), (4, "duplicate_in_file")}
    assert (
        _count(
            session_factory,
            company_alpha,
            "SELECT count(*) FROM products WHERE sku = :sku",
            sku=f"{prefix}-NEW",
        )
        == 0
    )


def test_missing_required_column_is_reported_on_the_header(
    api_client: TestClient, alpha_token: str
) -> None:
    response = _upload(api_client, alpha_token, "products", b"sku;nom\nA-1;Article\n")

    assert response.status_code == 422
    errors = response.json()["error"]["details"]["errors"]
    assert errors == [
        {
            "row": 1,
            "column": "Prix de vente",
            "code": "missing_column",
            "message": "Colonne absente.",
        }
    ]


# ------------------------------------------------------- formats de fichier


def test_windows_1252_csv_with_french_headers_and_numbers(
    api_client: TestClient,
    alpha_token: str,
    company_alpha: DemoCompany,
    session_factory: sessionmaker[Any],
) -> None:
    marker = uuid4().hex[:6]
    categories = f"Nom;Parent\r\nBoissons {marker};\r\nJus {marker};Boissons {marker}\r\n"
    assert (
        _upload(api_client, alpha_token, "categories", categories.encode("cp1252")).status_code
        == 201
    )

    products = (
        "SKU;Nom;Prix de vente;Catégorie;Prix d'achat\r\n"
        f"JUS-{marker};Jus de bissap;1 250,50;Jus {marker};800\r\n"
    ).encode("cp1252")
    response = _upload(api_client, alpha_token, "products", products)

    assert response.status_code == 201, response.text
    parent_of_jus = read_scalar_in_tenant(
        session_factory,
        company_alpha.company_id,
        "SELECT p.name FROM categories c JOIN categories p ON p.id = c.parent_id "
        "WHERE c.name = :name",
        {"name": f"Jus {marker}"},
    )
    assert parent_of_jus == f"Boissons {marker}"
    price = read_scalar_in_tenant(
        session_factory,
        company_alpha.company_id,
        "SELECT pp.unit_price FROM product_prices pp JOIN products p ON p.id = pp.product_id "
        "WHERE p.sku = :sku",
        {"sku": f"JUS-{marker}".upper()},
    )
    assert str(price) == "1250.50"


def test_xlsx_customers_import(
    api_client: TestClient,
    alpha_token: str,
    company_alpha: DemoCompany,
    session_factory: sessionmaker[Any],
) -> None:
    marker = uuid4().hex[:6]
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["Nom", "Téléphone", "Email", "Plafond crédit"])
    sheet.append([f"Awa {marker}", "+221 77 000 00 00", f"awa{marker}@example.com", 150000])
    sheet.append([f"Moussa {marker}", None, None, 0.1])
    buffer = io.BytesIO()
    workbook.save(buffer)

    response = _upload(
        api_client, alpha_token, "customers", buffer.getvalue(), filename="clients.xlsx"
    )

    assert response.status_code == 201, response.text
    assert response.json()["created"] == 2
    limit = read_scalar_in_tenant(
        session_factory,
        company_alpha.company_id,
        "SELECT credit_limit FROM customers WHERE name = :name",
        {"name": f"Moussa {marker}"},
    )
    # 0.1 lu dans Excel reste 0.10 : aucun flottant approximatif.
    assert str(limit) == "0.10"


def test_unreadable_file_is_refused(api_client: TestClient, alpha_token: str) -> None:
    response = _upload(api_client, alpha_token, "customers", b"PK\x03\x04pas-un-classeur")

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "import_file_unreadable"


# ------------------------------------------------------------------ idempotence


def test_replay_with_the_same_key_creates_nothing_more(
    api_client: TestClient,
    alpha_token: str,
    company_alpha: DemoCompany,
    session_factory: sessionmaker[Any],
) -> None:
    prefix = f"IDEM{uuid4().hex[:6].upper()}"
    key = uuid4().hex
    content = _products_csv(prefix, 3)

    first = _upload(api_client, alpha_token, "products", content, key=key)
    second = _upload(api_client, alpha_token, "products", content, key=key)
    other_file = _upload(api_client, alpha_token, "products", _products_csv(prefix, 4), key=key)

    assert first.status_code == second.status_code == 201
    assert second.json() == first.json()
    assert other_file.status_code == 422
    assert other_file.json()["error"]["code"] == "idempotency_key_reused"
    assert _state(session_factory, company_alpha, prefix)["products"] == 3


# ------------------------------------------------------ droits et modèles


def test_salesperson_cannot_import(
    api_client: TestClient, company_alpha: DemoCompany, alpha_token: str
) -> None:
    from tests.conftest import create_user_with_role

    token = create_user_with_role(
        api_client,
        alpha_token,
        company_alpha,
        role="vendeur",
        email=f"vendeur-{uuid4().hex[:6]}@societe-alpha.example.com",
    )
    response = _upload(api_client, token, "customers", b"nom\nX\n")

    assert response.status_code == 403


def test_template_lists_the_expected_columns(api_client: TestClient, alpha_token: str) -> None:
    response = api_client.get(
        "/api/v1/imports/products/template", headers=auth_headers(alpha_token)
    )

    assert response.status_code == 200
    header = response.content.decode("utf-8-sig").splitlines()[0]
    assert header.startswith("Référence;Nom;Prix de vente;Catégorie")


def test_exported_referential_reimports_as_is(
    api_client: TestClient, alpha_token: str, company_alpha: DemoCompany
) -> None:
    exported = api_client.get(
        "/api/v1/exports/referentials/suppliers",
        params={"format": "xlsx"},
        headers=auth_headers(alpha_token),
    )

    assert exported.status_code == 200, exported.text
    assert "attachment" in exported.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(exported.content)).worksheets[0]
    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert headers == ["Nom", "Code", "Téléphone", "Délai de paiement (jours)"]

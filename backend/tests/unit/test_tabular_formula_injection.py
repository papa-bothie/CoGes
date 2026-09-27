"""Neutralisation des formules à l'écriture des tableaux (ADR 0045)."""

from __future__ import annotations

import csv
import io
import zipfile
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from app.modules.importexport.tabular import (
    neutralise_formula,
    read_table,
    restore_formula_text,
    write_csv,
    write_xlsx,
)

DANGEROUS = [
    '=HYPERLINK("http://evil.example","x")',
    "+221 77 000 00 00",
    "-2+3",
    "@SUM(A1:A2)",
    "\tcmd",
    "\rcmd",
]
HARMLESS = ["Awa Diop", "", "'déjà cité", "12 rue du marché"]


@pytest.mark.parametrize("text", DANGEROUS)
def test_dangerous_text_is_prefixed_then_restored(text: str) -> None:
    assert neutralise_formula(text) == "'" + text
    assert restore_formula_text(neutralise_formula(text)) == text


@pytest.mark.parametrize("text", HARMLESS)
def test_harmless_text_is_left_alone(text: str) -> None:
    assert neutralise_formula(text) == text
    assert restore_formula_text(text) == text


def test_csv_prefixes_text_but_keeps_numbers() -> None:
    content = write_csv(["nom", "montant"], [[text, Decimal("-1500.50")] for text in DANGEROUS])

    rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig")), delimiter=";"))
    assert [row[0] for row in rows[1:]] == ["'" + text for text in DANGEROUS]
    # Un vrai nombre négatif reste un nombre : aucune apostrophe.
    assert {row[1] for row in rows[1:]} == {"-1500,50"}


def test_xlsx_has_no_formula_cell_and_shows_the_same_text() -> None:
    content = write_xlsx(
        ["nom", "montant", "quantite"],
        [[text, Decimal("-1500.50"), 3] for text in DANGEROUS],
        title="clients",
    )

    sheet_xml = zipfile.ZipFile(io.BytesIO(content)).read("xl/worksheets/sheet1.xml")
    assert b"<f>" not in sheet_xml
    sheet = load_workbook(io.BytesIO(content)).worksheets[0]
    body = list(sheet.iter_rows(min_row=2))
    assert all(cell.data_type != "f" for row in body for cell in row)
    # Tout parseur XML rend « \r » en « \n » (fin de ligne normalisée) : propre au format.
    assert [row[0].value for row in body] == [text.replace("\r", "\n") for text in DANGEROUS]
    assert all(row[0].data_type == "s" for row in body)
    # Les nombres restent des nombres, pas du texte.
    assert all(row[1].data_type == "n" and row[2].data_type == "n" for row in body)
    assert {row[1].value for row in body} == {-1500.5}


@pytest.mark.parametrize("render", ["csv", "xlsx"])
def test_round_trip_gives_back_the_original_text(render: str) -> None:
    values = [text for text in DANGEROUS if text.strip() == text] + HARMLESS[:2]
    rows = [[text] for text in values]
    content = write_csv(["nom"], rows) if render == "csv" else write_xlsx(["nom"], rows, title="t")

    table = read_table(content, max_rows=100)
    assert [row.cells["nom"] for row in table.rows] == [text for text in values if text]

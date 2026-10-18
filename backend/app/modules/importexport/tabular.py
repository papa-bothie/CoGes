"""Lecture et écriture de tableaux CSV et Excel (.xlsx), sans règle métier.

Un fichier de gérant ouest-africain arrive rarement « propre » : exporté d'Excel en
français, il est souvent en Windows-1252, séparé par des points-virgules, avec des
nombres écrits `1 234,50`. La lecture tolère tout cela ; la validation, elle, est
faite ensuite par les schémas Pydantic des modules, cellule par cellule.

À l'écriture, le CSV suit la même convention (UTF-8 avec BOM, `;`, virgule
décimale) : il s'ouvre tel quel dans un Excel réglé en français, et il se réimporte.

Aucune cellule exportée n'est une formule (ADR 0045) : un nom de client
`=HYPERLINK(…)` saisi par un tiers ne doit pas s'exécuter chez le gérant qui ouvre
l'export. En CSV, un texte qui commence par `=`, `+`, `-`, `@`, une tabulation ou un
retour chariot est préfixé d'une apostrophe ; en .xlsx, tout texte est écrit en type
chaîne. Les nombres restent des nombres. La lecture retire l'apostrophe ajoutée :
l'aller-retour export → import rend la valeur d'origine.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any

#: Signature d'une archive ZIP : un .xlsx en est une.
_XLSX_MAGIC = b"PK\x03\x04"

#: Bornes d'un fichier importé (ADR 0088). Un .xlsx est un zip : quelques Ko peuvent se
#: décompresser en Go. On lit la table des matières **avant** d'ouvrir le classeur.
MAX_UNCOMPRESSED_BYTES = 30 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_COLUMNS = 100
MAX_CELL_CHARS = 2000

#: Petits mots ignorés dans un en-tête.
_STOP_WORDS = frozenset({"de", "d", "du", "des", "la", "le", "l", "en"})

#: Séparateurs acceptés en CSV, dans l'ordre de préférence en cas d'égalité.
_DELIMITERS = (";", ",", "\t")


#: Premiers caractères qui font d'une cellule une formule (ou un début de formule) dans
#: un tableur (ADR 0045).
_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")

#: Préfixe qui neutralise ces caractères en CSV.
_TEXT_PREFIX = "'"


def is_formula_like(text: str) -> bool:
    """Vrai si un tableur lirait ce texte comme une formule."""
    return text.startswith(_FORMULA_TRIGGERS)


def neutralise_formula(text: str) -> str:
    """`=1+1` → `'=1+1` ; un texte inoffensif est rendu tel quel."""
    return _TEXT_PREFIX + text if is_formula_like(text) else text


def restore_formula_text(text: str) -> str:
    """Inverse de `neutralise_formula` : `'=1+1` → `=1+1`, et rien d'autre ne change."""
    if text.startswith(_TEXT_PREFIX) and is_formula_like(text[len(_TEXT_PREFIX) :]):
        return text[len(_TEXT_PREFIX) :]
    return text


class UnreadableFileError(ValueError):
    """Fichier vide, illisible, ou d'un format qui n'est ni CSV ni .xlsx."""


@dataclass(frozen=True, slots=True)
class Row:
    """Une ligne de données, avec son numéro dans le fichier (l'en-tête est la ligne 1)."""

    number: int
    cells: dict[str, str]


@dataclass(frozen=True, slots=True)
class Table:
    headers: list[str]
    rows: list[Row]


def plain(value: str) -> str:
    """Minuscules sans accents : `Pièce ` → `piece`."""
    decomposed = unicodedata.normalize("NFKD", value.strip().casefold())
    return "".join(char for char in decomposed if not unicodedata.combining(char))


def normalise_header(value: str) -> str:
    """`Prix de vente ` → `prix_vente` ; `Prix d'achat` → `prix_achat` ; `Catégorie` → `categorie`.

    Casse, accents, ponctuation et petits mots (« de », « du », « d' »…) sont ignorés :
    un en-tête retapé à la main dans Excel se reconnaît quand même.
    """
    words = [word for word in re.findall(r"[a-z0-9]+", plain(value)) if word not in _STOP_WORDS]
    return "_".join(words)


def _cell_text(value: Any) -> str:
    """Texte d'une cellule Excel, sans passer par un flottant approximatif."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, float):
        # `repr` rend la plus courte écriture qui redonne le même flottant :
        # 0.1 reste « 0.1 », pas 0.1000000000000000055…
        text = repr(value)
        return text[:-2] if text.endswith(".0") else text
    if isinstance(value, datetime | date):
        return value.isoformat()
    return str(value).strip()


def _check_archive(content: bytes) -> None:
    """Refuse une bombe zip sans rien décompresser : tailles lues dans l'en-tête."""
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
    except zipfile.BadZipFile as exc:
        raise UnreadableFileError("Le fichier Excel est illisible.") from exc
    unpacked = sum(entry.file_size for entry in entries)
    packed = sum(entry.compress_size for entry in entries) or 1
    if unpacked > MAX_UNCOMPRESSED_BYTES:
        raise UnreadableFileError(
            f"Le fichier Excel est trop gros une fois décompressé ({unpacked // 1_048_576} Mo, "
            f"{MAX_UNCOMPRESSED_BYTES // 1_048_576} au plus)."
        )
    if unpacked > 1_048_576 and unpacked / packed > MAX_COMPRESSION_RATIO:
        raise UnreadableFileError("Le fichier Excel est compressé de façon anormale : refusé.")


def _read_xlsx(content: bytes, *, max_rows: int) -> list[list[str]]:
    from openpyxl import load_workbook  # noqa: PLC0415 - chargé seulement pour un .xlsx

    _check_archive(content)
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # zip corrompu, XML invalide…
        raise UnreadableFileError("Le fichier Excel est illisible.") from exc
    try:
        sheet = workbook.worksheets[0]
        lines: list[list[str]] = []
        # Une colonne de plus que la borne suffit à refuser, sans lire le reste de la ligne.
        for values in sheet.iter_rows(values_only=True, max_col=MAX_COLUMNS + 1):
            cells = [_cell_text(value) for value in values]
            # `max_col` complète la ligne de cellules vides : elles ne comptent pas.
            while cells and not cells[-1]:
                cells.pop()
            lines.append(cells)
            if len(lines) > max_rows + 1:
                break
        return lines
    finally:
        workbook.close()


def _decode(content: bytes) -> str:
    try:
        return content.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    # Excel français enregistre en Windows-1252 : « é » y est l'octet 0xE9. Un octet que
    # ce codage ne connaît pas est signalé, jamais remplacé par « � » (ADR 0088).
    try:
        return content.decode("cp1252")
    except UnicodeDecodeError as exc:
        raise UnreadableFileError(
            "Le fichier n'est ni en UTF-8 ni en Windows-1252 : enregistrez-le depuis Excel "
            "en « CSV UTF-8 », puis réessayez."
        ) from exc


def _delimiter(first_line: str) -> str:
    counts = {delimiter: first_line.count(delimiter) for delimiter in _DELIMITERS}
    best = max(_DELIMITERS, key=lambda delimiter: counts[delimiter])
    return best if counts[best] else ";"


def _read_csv(content: bytes, *, max_rows: int) -> list[list[str]]:
    text = _decode(content)
    first_line = text.split("\n", 1)[0]
    reader = csv.reader(io.StringIO(text), delimiter=_delimiter(first_line))
    lines: list[list[str]] = []
    try:
        for values in reader:
            lines.append([value.strip() for value in values])
            if len(lines) > max_rows + 1:
                break
    except csv.Error as exc:
        raise UnreadableFileError("Le fichier CSV est mal formé.") from exc
    return lines


def _check_shape(lines: list[list[str]]) -> None:
    """Colonnes et longueur de cellule bornées (ADR 0088)."""
    for number, values in enumerate(lines, start=1):
        if len(values) > MAX_COLUMNS:
            raise UnreadableFileError(
                f"Le fichier a trop de colonnes ({MAX_COLUMNS} au plus, ligne {number})."
            )
        if any(len(value) > MAX_CELL_CHARS for value in values):
            raise UnreadableFileError(
                f"Une cellule dépasse {MAX_CELL_CHARS} caractères (ligne {number})."
            )


def read_table(content: bytes, *, max_rows: int) -> Table:
    """Lit un CSV ou un .xlsx ; lignes entièrement vides ignorées.

    Lit au plus `max_rows + 1` lignes de données : l'appelant sait ainsi qu'il y en a
    trop sans que le fichier entier soit chargé.
    """
    if not content.strip():
        raise UnreadableFileError("Le fichier est vide.")
    lines = (
        _read_xlsx(content, max_rows=max_rows)
        if content.startswith(_XLSX_MAGIC)
        else _read_csv(content, max_rows=max_rows)
    )
    if not lines:
        raise UnreadableFileError("Le fichier est vide.")
    _check_shape(lines)
    headers = [normalise_header(value) for value in lines[0]]
    seen: set[str] = set()
    for position, header in enumerate(headers):
        if header and header in seen:
            raise UnreadableFileError(
                f"La colonne « {lines[0][position].strip()} » apparaît deux fois : "
                "gardez-en une seule."
            )
        seen.add(header)
    rows: list[Row] = []
    for index, values in enumerate(lines[1:], start=2):
        if not any(values):
            continue
        cells = {
            header: (restore_formula_text(values[position]) if position < len(values) else "")
            for position, header in enumerate(headers)
            if header
        }
        rows.append(Row(number=index, cells=cells))
    return Table(headers=[header for header in headers if header], rows=rows)


# ------------------------------------------------------------------- écriture


def _csv_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, Decimal):
        # Virgule décimale : le séparateur de champ est le point-virgule.
        return format(value, "f").replace(".", ",")
    if isinstance(value, str):
        return neutralise_formula(value)
    return str(value)


def write_csv(headers: Sequence[str], rows: Iterable[Sequence[Any]]) -> bytes:
    """CSV UTF-8 avec BOM, séparé par `;` : s'ouvre directement dans Excel en français."""
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    writer.writerow([_csv_value(header) for header in headers])
    for row in rows:
        writer.writerow([_csv_value(value) for value in row])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


def write_xlsx(headers: Sequence[str], rows: Iterable[Sequence[Any]], *, title: str) -> bytes:
    """Classeur à une feuille ; les montants restent des nombres (Decimal), pas du texte.

    Tout texte est une cellule de type chaîne : openpyxl ferait sinon d'un `=…` une
    formule. Le style « quotePrefix » garde en texte une cellule `=…`, `+…`, `-…`
    même après une modification dans Excel, sans apostrophe visible.
    """
    from openpyxl import Workbook  # noqa: PLC0415
    from openpyxl.cell import WriteOnlyCell  # noqa: PLC0415

    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet(title=title[:31])

    def cell(value: Any) -> Any:
        if isinstance(value, bool):
            value = "oui" if value else "non"
        if not isinstance(value, str):
            return value
        text_cell = WriteOnlyCell(sheet, value)
        text_cell.data_type = "s"
        if is_formula_like(value):
            text_cell.quotePrefix = True
        return text_cell

    sheet.append([cell(header) for header in headers])
    for row in rows:
        sheet.append([cell(value) for value in row])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
